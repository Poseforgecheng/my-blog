from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPORT_TITLE = "AI 开发一个人网站"
REPORT_SUBTITLE = "基于个人网站实践的 AI 协同开发方法总结"
REPORT_DATE = "2026年3月17日"
AUTHOR = "ChengzhuoSir"
PROJECT_NAME = "个人网站建设与 AI 协同开发总结"


def _set_run_font(run, ascii_font: str = "Times New Roman", east_asia_font: str = "宋体", size: int = 12, bold: bool = False):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)
    run.font.size = Pt(size)
    run.bold = bold


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, center: bool = False, size: int = 11) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    _set_run_font(run, east_asia_font="宋体", size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    _set_run_font(run, east_asia_font="宋体", size=9)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.color.rgb = RGBColor(31, 78, 121)

    document.core_properties.title = REPORT_TITLE
    document.core_properties.author = AUTHOR
    document.core_properties.subject = REPORT_SUBTITLE
    document.core_properties.comments = "Generated with python-docx"


def _configure_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    header_para = header.paragraphs[0]
    header_para.text = ""
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run(REPORT_TITLE)
    _set_run_font(run, east_asia_font="宋体", size=9)
    run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    footer_para.text = ""
    _add_page_number(footer_para)


def _add_paragraph(document: Document, text: str, *, indent: bool = True, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        _set_run_font(lead, east_asia_font="宋体", size=12, bold=True)
        body = paragraph.add_run(text[len(bold_lead):])
        _set_run_font(body, east_asia_font="宋体", size=12)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run, east_asia_font="宋体", size=12)


def _add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = paragraph.add_run(item)
        _set_run_font(run, east_asia_font="宋体", size=12)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], title: str | None = None) -> None:
    if title:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        _set_run_font(run, east_asia_font="黑体", size=12, bold=True)
        paragraph.paragraph_format.space_after = Pt(6)

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        _set_cell_text(header_cells[idx], header, bold=True, center=True, size=11)
        _shade_cell(header_cells[idx], "D9EAF7")

    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            _set_cell_text(cells[idx], value, center=False, size=10 if len(value) > 45 else 11)

    document.add_paragraph()


def _add_cover_page(document: Document) -> None:
    for _ in range(6):
        document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(REPORT_TITLE)
    _set_run_font(title_run, east_asia_font="黑体", size=24, bold=True)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(REPORT_SUBTITLE)
    _set_run_font(subtitle_run, east_asia_font="宋体", size=14)

    for _ in range(8):
        document.add_paragraph()

    meta = [
        f"汇报对象：领导汇报",
        f"项目名称：{PROJECT_NAME}",
        f"作者：{AUTHOR}",
        f"日期：{REPORT_DATE}",
    ]
    for line in meta:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(line)
        _set_run_font(run, east_asia_font="宋体", size=14)

    document.add_page_break()


def _add_abstract(document: Document) -> None:
    document.add_heading("摘要", level=1)
    _add_paragraph(
        document,
        "本报告围绕“AI 如何参与个人网站开发”这一主题，结合当前个人网站项目的真实实践，对 AI 在前端设计、内容整理、格式规范、问题修复、发布验证以及后续能力扩展中的作用进行了系统梳理。项目当前采用 Jekyll 与 GitHub Pages 构建静态网站，已完成首页、博客、关于页、游戏页等基础模块建设，并通过本地化的 Mermaid、MathJax 与前端脚本能力提升内容表达与阅读体验。",
    )
    _add_paragraph(
        document,
        "在建设过程中，AI 并非仅承担代码生成工具的角色，而是作为协同开发助手贯穿需求整理、页面结构设计、文案优化、格式修订、问题排查与经验沉淀多个环节。实践表明，AI 能显著缩短从想法到页面、从零散笔记到可发布内容、从问题发现到规范固化的路径，使个人网站建设逐步形成低门槛、可持续、可迭代的工程化流程。",
    )
    _add_paragraph(
        document,
        "基于当前静态网站形态，报告还提出了后续接入 AI 后端/API 的扩展方案，包括智能问答、内容摘要、自动标签、搜索增强与写作助手等能力，说明该项目可以从内容展示型网站逐步演进为具备智能服务能力的个人知识门户。整体上，本项目验证了 AI 协同开发在网站建设场景中的现实价值，也为后续平台化演进提供了方法基础。",
    )

    document.add_heading("关键词", level=2)
    _add_paragraph(document, "AI协同开发；个人网站；前端设计；内容管理；格式规范；后端扩展", indent=False)
    document.add_page_break()


def _add_toc(document: Document) -> None:
    document.add_heading("目录", level=1)
    toc_items = [
        "1. 项目背景与目标",
        "2. 网站现状与建设成果",
        "3. AI 参与开发的总体思路",
        "4. AI 如何参与前端开发",
        "5. AI 如何参与内容整理与知识沉淀",
        "6. AI 如何参与格式规范与文档修改",
        "7. AI 如何辅助测试、修复与发布",
        "8. 可扩展方案：如果接入 AI 后端/API，应该怎么设计",
        "9. 实施价值与后续建议",
        "附录 A. 当前网站关键实现要点",
        "附录 B. AI 协同开发提示语模板",
    ]
    for item in toc_items:
        _add_paragraph(document, item, indent=False)
    document.add_page_break()


def _add_main_sections(document: Document) -> None:
    document.add_heading("1. 项目背景与目标", level=1)
    _add_paragraph(document, "个人网站建设的直接目标，是形成一个可持续维护的个人表达与知识沉淀平台。与一次性展示页面不同，个人网站需要同时兼顾结构清晰、内容可扩展、风格统一和维护成本可控等要求。对于持续输出技术内容的人而言，网站既是外部展示窗口，也是内部工作记录、项目复盘与方法论沉淀的载体。")
    _add_paragraph(document, "传统网站开发方式通常需要在页面设计、代码编写、内容整理、格式修订和发布维护之间频繁切换，个人开发者在时间和精力有限的情况下，很容易陷入“功能做出来了，但内容难以持续积累”“页面能打开，但格式难以统一”“文章已经写完，但上线效率较低”等问题。")
    _add_paragraph(document, "因此，本项目引入 AI 的核心目的，不是简单用工具替代编写代码，而是将 AI 作为协同开发助手，参与需求拆解、页面生成、表达优化、规范检查和问题修复等环节，帮助个人网站从零散开发转向流程化建设。这样的网站建设方式，更接近“以 AI 提升个人工程产能”的实践路径。")

    document.add_heading("2. 网站现状与建设成果", level=1)
    _add_paragraph(document, "从当前项目状态看，网站已经具备较完整的静态站点形态。底层采用 Jekyll 构建与 GitHub Pages 部署，既满足了低成本上线的要求，也方便后续持续维护。站点配置中已经明确了标题、描述、时区、基础路径、Markdown 解析方式和扩展插件，说明项目已经具备稳定的内容发布基础。")
    _add_paragraph(document, "页面层面，当前站点已包含首页、博客页、关于页和游戏页等核心模块。首页承担个人定位、内容入口和风格展示作用；博客页用于按时间组织技术文章；关于页用于补充个人简介与站点定位；游戏页作为实验性内容入口，扩展了网站的展示边界。整体上，网站已经不再是单页原型，而是具备清晰导航、模块组织和内容承载能力的个人站点。")
    _add_paragraph(document, "内容呈现层面，网站已经引入本地 Mermaid 与 MathJax 资源，用于支持流程图和公式渲染；同时通过前端脚本实现内容渐显、复制按钮、图表渲染等交互细节，说明项目已经开始从“能看”迈向“可读、可用、可维护”。这些真实实现为本报告关于 AI 协同开发的分析提供了直接案例基础。")
    _add_table(
        document,
        ["模块", "当前状态", "说明"],
        [
            ["首页", "已完成", "用于个人定位展示、最新文章入口和视觉风格表达。"],
            ["博客系统", "已完成", "支持技术文章的时间序组织与持续更新。"],
            ["关于页面", "已完成", "说明作者背景、站点定位与更新方向。"],
            ["游戏页面", "已完成", "作为扩展展示模块，体现实验性与趣味性。"],
            ["图表与公式", "已完成", "通过本地 Mermaid 与 MathJax 支持技术内容表达。"],
        ],
        "表 1 当前网站建设成果概览",
    )

    document.add_heading("3. AI 参与开发的总体思路", level=1)
    _add_paragraph(document, "在本项目中，AI 的定位不是独立完成所有开发工作的“自动开发者”，而是围绕个人网站建设各环节提供高频辅助的“协同开发助手”。其主要价值在于：帮助快速形成初稿、减少重复劳动、统一表达风格、提高问题定位效率，并将零散经验沉淀为可复用规则。")
    _add_paragraph(document, "从流程角度看，AI 的参与形成了一个比较清晰的闭环：先根据网站目标和受众整理需求，再辅助生成页面结构和基础文案；随后参与内容整理、段落优化和格式统一；在发布前后，继续辅助检查路径、配置、渲染和兼容性问题；最后将修改经验回收为规范文档，为下一轮更新提供模板。")
    _add_paragraph(document, "这种闭环式使用方式的意义在于，AI 不再只在“写代码”这一点上发挥作用，而是深入到网站开发的完整生命周期。这样既能保证开发效率，也能让个人网站的长期维护成本逐步下降。")
    _add_table(
        document,
        ["环节", "AI 主要作用", "人工把关重点", "典型产出"],
        [
            ["需求整理", "整理目标、提炼结构、形成章节或页面框架", "确认方向是否符合用途", "页面草案、文档大纲"],
            ["页面开发", "生成结构代码、优化文案、补齐样式与交互", "确认视觉与功能符合预期", "首页、卡片、导航、交互脚本"],
            ["内容整理", "提炼笔记重点、生成摘要、统一表述", "校验技术准确性与表达尺度", "博客正文、摘要、说明文档"],
            ["格式规范", "修正文档层级、公式、表格、流程图写法", "确认规范是否适合站点长期维护", "Markdown 规则、排版规范"],
            ["修复发布", "辅助定位路径、渲染、构建与部署问题", "验证修复是否真实生效", "修复方案、检查清单、复盘记录"],
        ],
        "表 2 AI 参与网站开发的闭环作用表",
    )

    document.add_heading("4. AI 如何参与前端开发", level=1)
    _add_paragraph(document, "前端开发是 AI 最容易产生直接价值的环节之一。对于个人网站而言，页面通常需要在较短时间内完成结构搭建、风格统一和文案填充，而 AI 擅长根据自然语言描述迅速给出页面区块、标题、按钮、说明文案和视觉布局建议，从而缩短“想法到首版页面”的周期。")
    _add_paragraph(document, "结合本项目实际，首页的英雄区、关注标签、卡片式内容区和导航结构，都非常适合由 AI 先给出结构化初稿，再由人工进行调整。这样做的好处是可以快速确定页面层级和信息重点，避免个人开发者在布局起步阶段耗费过多时间。")
    _add_paragraph(document, "除了页面骨架，AI 在样式层面也能够承担辅助角色。例如，可以先给出一套符合站点定位的颜色、间距、按钮样式、动效建议，再由开发者结合实际审美和仓库现有风格进行筛选。当前站点已经包含统一的 CSS 变量、背景视觉、卡片样式和按钮形态，这类统一性工作本身就非常适合在 AI 协助下完成。")
    _add_paragraph(document, "在交互层面，AI 也适合辅助补齐一些标准化前端能力，例如渐显动效、Mermaid 渲染、代码复制按钮和数学公式延迟渲染逻辑。这些逻辑本身并不复杂，但细节较多、易出现兼容性遗漏。AI 可以帮助快速组织脚本结构，并提示需要覆盖的边界条件。")
    _add_bullets(
        document,
        [
            "用于快速产出首页、博客页、专题页等页面的初步结构。",
            "用于补全文案、按钮语义、模块标题和引导说明，使页面更完整。",
            "用于提出样式统一建议，减少视觉碎片化问题。",
            "用于生成标准交互脚本，补齐页面体验细节。",
        ],
    )
    _add_paragraph(document, "需要强调的是，AI 在前端开发中的最佳用法并不是直接全盘接受其输出，而是将其视为高效率的“结构建议器”和“草稿生成器”。最终页面是否符合个人定位、审美风格和内容节奏，仍然需要开发者做最后判断。")

    document.add_heading("5. AI 如何参与内容整理与知识沉淀", level=1)
    _add_paragraph(document, "个人网站的可持续价值，本质上来自内容，而不是页面本身。很多技术人员的实际问题并不是“不会写网页”，而是“工作记录零散、知识没有结构化、写过的内容难以二次利用”。在这种场景下，AI 的价值非常明显：它能够把分散的笔记、会议记录、调试心得和项目片段整理成可发布的文章或说明文档。")
    _add_paragraph(document, "当前仓库中已经存在多篇技术文章和两份规则类说明文档，这说明网站并不只是展示页面，而是在逐渐形成知识资产。AI 可以帮助从原始笔记中提炼出主题、摘要、章节层级与段落过渡，使文章从“作者自己看得懂”转变为“外部读者也能理解”。")
    _add_paragraph(document, "进一步看，AI 还可以承担内容归档与复用工作。例如，对文章标题进行统一命名，对摘要进行标准化，对同类笔记进行合并，对旧文进行结构重写，对长期系列文章维持一致的表达风格。这样不仅提升阅读体验，也能降低个人网站长期维护时的整理成本。")
    _add_paragraph(document, "对于面向技术积累的网站来说，AI 最重要的作用之一，就是把碎片化经验转化为可搜索、可回顾、可复用的知识资产。这一点对个人品牌建设、项目复盘和后续扩展都非常关键。")

    document.add_heading("6. AI 如何参与格式规范与文档修改", level=1)
    _add_paragraph(document, "网站内容一旦开始积累，格式规范就会变得重要。尤其是技术文章中常见的 Markdown 标题层级、代码块、公式、Mermaid 图表和表格写法，如果缺乏统一规则，后续维护会迅速失控。本项目已经有较明确的 Markdown 写作规则和站点构建注意事项，这为 AI 参与格式治理提供了直接抓手。")
    _add_paragraph(document, "在这一环节，AI 的优势不在于“创造新内容”，而在于“根据既有规则批量修正内容”。例如，AI 可以检查公式中不安全的竖线写法、修正表格前后空行、统一代码块语言标记、调整章节层级连续性，或者把不适合网页渲染的段落结构重写为更稳定的形式。")
    _add_paragraph(document, "对网站维护者而言，这种能力非常实用。因为很多问题并不是内容错误，而是格式不统一导致的显示异常、渲染失败或页面风格不一致。将这些经验沉淀为规则文档之后，AI 就能够围绕规则执行批量化修订，显著降低人工逐条检查的压力。")
    _add_table(
        document,
        ["规范场景", "AI 可执行动作", "实际收益"],
        [
            ["标题层级", "统一一级、二级、三级标题连续性", "增强结构清晰度，便于长期维护"],
            ["公式书写", "修正不安全符号与不稳定排版", "减少 MathJax 渲染问题"],
            ["Mermaid 图表", "补足代码块边界与图表文本规范", "提升流程图渲染稳定性"],
            ["表格与代码块", "补空行、补语言标记、修正格式粘连", "减少构建后显示异常"],
            ["文档重写", "把口语化或零散笔记重写为正式表达", "提升对外展示质量"],
        ],
        "表 3 AI 参与格式治理的典型场景",
    )

    document.add_heading("7. AI 如何辅助测试、修复与发布", level=1)
    _add_paragraph(document, "网站开发并不止于页面完成，真正影响交付质量的，往往是最后阶段的测试、修复和发布。对于静态网站而言，常见问题包括资源路径错误、基础路径配置不一致、本地渲染正常但线上资源失效、Mermaid 或 MathJax 未加载、CSS 与 JS 更新后缓存未刷新等。")
    _add_paragraph(document, "AI 在这一阶段的价值，主要体现在帮助开发者快速缩小排查范围。只要能够提供出错现象、构建命令、配置片段或控制台信息，AI 通常可以迅速从资源路径、配置项、构建顺序、依赖加载和部署环境等角度给出高概率原因，并组织成排查清单。")
    _add_paragraph(document, "更重要的是，AI 可以把一次次修复过程沉淀为规范。例如，本项目中已经形成了关于 Jekyll 构建、资源命名、Mermaid 渲染、MathJax 类型设置和 Markdown 规则的说明文档。这意味着问题处理不再只停留在“修好了”，而是逐步固化成“以后怎么避免再次出现”。")
    _add_paragraph(document, "对于个人网站这种持续维护的项目而言，这种从修复到规范化的转变非常关键。AI 不仅帮助解决当前问题，也帮助减少下一次问题。")

    document.add_heading("8. 可扩展方案：如果接入 AI 后端/API，应该怎么设计", level=1)
    _add_paragraph(document, "虽然当前网站采用的是静态站点方案，但从演进角度看，它完全可以进一步接入 AI 后端/API，形成“静态内容承载 + 动态智能服务增强”的混合架构。这一扩展不应理解为推翻现有方案，而应理解为在保持低成本内容发布能力的前提下，为网站增加智能交互能力。")
    _add_paragraph(document, "第一层是能力层，建议优先考虑与个人网站价值最相关的能力，例如文章摘要、自动标签、站内智能问答、相关文章推荐、长文提炼和写作助手。这些能力既能提升读者体验，也能反向提升站长的内容管理效率。")
    _add_paragraph(document, "第二层是服务层，可设计为独立的后端 API，包括用户请求接入、鉴权与限流、内容服务、向量检索、模型调用、日志监控和缓存模块。若未来网站文章数量持续增长，还可以引入知识库切片与检索增强，让读者直接在站内围绕个人文章和笔记发起问答。")
    _add_paragraph(document, "第三层是应用层，也就是当前网站前端的扩展入口。前端页面可以增加问答组件、摘要按钮、推荐卡片、自动标签和辅助写作界面，以 API 调用方式连接后端模型服务。这使得网站从“内容展示平台”演进为“内容与智能服务结合的平台”。")
    _add_table(
        document,
        ["层次", "模块", "说明"],
        [
            ["应用层", "问答组件、摘要按钮、推荐卡片、写作助手界面", "直接面向读者与站长提供交互能力"],
            ["服务层", "API 网关、鉴权限流、内容服务、检索服务、模型编排", "负责请求路由、能力封装和运行治理"],
            ["数据层", "文章正文、标签、索引、向量库、日志数据", "支撑内容检索、推荐和效果分析"],
            ["模型层", "摘要模型、问答模型、标签模型、文本润色模型", "提供可插拔的智能能力来源"],
        ],
        "表 4 AI 后端/API 扩展架构示意",
    )
    _add_paragraph(document, "从实施策略看，建议后续扩展遵循“小步快跑、能力分层”的原则。可以先做不改变当前站点架构的小能力，例如离线摘要、自动标签和文章相关推荐；再逐步加入在线问答、检索增强和内容助手；最后视使用情况决定是否引入更完整的后端服务治理体系。")

    document.add_heading("9. 实施价值与后续建议", level=1)
    _add_paragraph(document, "综合来看，AI 协同开发方式在个人网站建设中至少体现了四方面价值。第一，是明显降低了页面搭建和内容整理的启动门槛，使个人开发者能够把更多精力放在判断和迭代上，而不是消耗在重复操作上。第二，是帮助内容生产与规范维护形成闭环，减少“写完即散落”的情况。")
    _add_paragraph(document, "第三，AI 让网站维护逐步具备工程化特征。需求可以先整理，页面可以先打样，问题可以先归因，规则可以先落文档，这使得个人网站不再只是零散创作，而更像一个可长期经营的小型产品。第四，若未来继续接入后端/API，当前做法还能自然迁移，形成从静态网站到智能知识门户的演进基础。")
    _add_paragraph(document, "建议下一阶段从三个方向推进：一是继续沉淀内容规范与发布流程，让 AI 更容易按规则协同工作；二是围绕重点主题建设专题文章与系列化栏目，提升网站内容资产密度；三是择机验证一个轻量 AI 服务能力，例如自动摘要或站内问答，为后续后端扩展提供真实使用反馈。")
    _add_paragraph(document, "总体而言，本项目已经证明，AI 在个人网站开发中最有价值的地方，并不是单点替代某项技术工作，而是在多个环节持续提高组织效率、表达质量和迭代速度。对个人开发者而言，这是一条现实可行、可持续演进的建设路径。")


def _add_appendices(document: Document) -> None:
    document.add_page_break()
    document.add_heading("附录 A. 当前网站关键实现要点", level=1)
    _add_paragraph(document, "为确保本报告建立在项目真实状态之上，以下内容概括了当前仓库可直接验证的关键实现点。这些实现点既是网站现状的证明，也是 AI 参与开发分析的案例来源。")
    _add_table(
        document,
        ["类别", "仓库中的体现", "对汇报的意义"],
        [
            ["站点框架", "Jekyll 配置、GitHub Pages 部署、固定 permalink 与插件配置", "说明网站已形成稳定的静态发布链路"],
            ["页面结构", "首页、博客、关于页、游戏页与默认布局", "说明站点具备可浏览的多页面形态"],
            ["前端表达", "统一 CSS 变量、卡片样式、背景视觉和按钮体系", "说明 AI 可参与页面风格统一与迭代"],
            ["脚本能力", "Mermaid 渲染、MathJax 调用、代码复制按钮、页面 reveal 动效", "说明 AI 可辅助标准交互能力补齐"],
            ["内容规范", "Markdown 写作规则与 Jekyll 构建注意事项文档", "说明 AI 可围绕规则执行格式治理"],
        ],
    )
    _add_paragraph(document, "这些实现点意味着当前项目已经具备“以真实项目支撑方法论总结”的条件。报告中关于 AI 的作用，并不是脱离项目的泛泛而谈，而是从当前站点建设过程反推出的协同开发方法。")

    document.add_heading("附录 B. AI 协同开发提示语模板", level=1)
    _add_paragraph(document, "为了提升后续使用效率，建议将一部分常用提示语模板固化下来，使 AI 在不同阶段更快进入合适角色。以下模板均可根据具体任务裁剪。")
    _add_table(
        document,
        ["场景", "建议提示语", "预期结果"],
        [
            ["页面设计", "请基于个人技术网站定位，给出首页结构、模块标题、按钮文案和视觉方向。", "快速形成页面草案"],
            ["内容整理", "请把下面的零散技术笔记整理成适合博客发布的文章结构，并补充摘要。", "形成可发布内容初稿"],
            ["格式修订", "请按既有 Markdown 规则检查本文，重点修正标题层级、公式、表格和代码块格式。", "减少格式类错误"],
            ["问题排查", "请根据以下构建日志和页面现象，分析可能的路径、配置或资源加载问题。", "快速形成排查清单"],
            ["后端设计", "请基于当前静态网站，给出接入 AI 摘要、问答和推荐能力的后端/API 扩展方案。", "形成演进设计草案"],
        ],
        "表 5 AI 协同开发常用提示语模板",
    )
    _add_paragraph(document, "将提示语模板化的意义，在于把“会不会用 AI”转化为“是否已经形成一套可复用的协作方式”。对于需要长期维护的网站项目，这种可复用性比单次生成结果更重要。")


def build_report(output_path: str | Path) -> dict[str, str | int]:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _style_document(document)
    _configure_header_footer(document.sections[0])
    _add_cover_page(document)
    _add_abstract(document)
    _add_toc(document)
    _add_main_sections(document)
    _add_appendices(document)

    for section in document.sections:
        _configure_header_footer(section)

    document.save(output)
    return {
        "output_path": str(output),
        "table_count": len(document.tables),
        "section_count": 11,
    }


def main() -> None:
    project_root = Path(__file__).resolve().parent.parent
    target = project_root / "output" / "doc" / "AI 开发一个人网站.docx"
    result = build_report(target)
    print(f"Generated: {result['output_path']}")
    print(f"Tables: {result['table_count']}")
    print(f"Sections: {result['section_count']}")


if __name__ == "__main__":
    main()
