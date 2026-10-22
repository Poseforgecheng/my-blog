import tempfile
import unittest
from pathlib import Path

from docx import Document

from tools.generate_ai_personal_website_report import build_report


class AiPersonalWebsiteReportTest(unittest.TestCase):
    def test_build_report_creates_expected_structure(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            target = Path(tmpdir) / "AI 开发一个人网站.docx"

            result = build_report(target)

            self.assertTrue(target.exists())
            self.assertEqual(result["output_path"], str(target))

            document = Document(target)
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            full_text = "\n".join(paragraphs)

            self.assertIn("AI 开发一个人网站", full_text)
            self.assertIn("项目背景与目标", full_text)
            self.assertIn("AI 如何参与前端开发", full_text)
            self.assertIn("可扩展方案：如果接入 AI 后端/API，应该怎么设计", full_text)
            self.assertGreaterEqual(len(document.tables), 3)


if __name__ == "__main__":
    unittest.main()
